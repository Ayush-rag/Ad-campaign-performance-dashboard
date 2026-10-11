"""Step 4 - Run the SQL KPI queries and turn the results into decisions.

Outputs:
  outputs/sql_results/*.csv   one CSV per SQL query
  outputs/ad_kpi_report.xlsx  formatted Excel report (one sheet per analysis)
  outputs/insights.md         key findings and budget recommendations
"""
import re
import sqlite3

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import (DB_PATH, EXCEL_REPORT, KPI_SQL, OUTPUT_DIR,
                    REALLOCATION_EFFICIENCY, TARGET_ROAS, roas_action)


def load_queries(path) -> dict[str, str]:
    """Split the .sql file into {query_name: sql} using the '-- name:' markers."""
    parts = re.split(r"^-- name:\s*(\w+)\s*$", path.read_text(), flags=re.MULTILINE)
    return {name: sql.strip() for name, sql in zip(parts[1::2], parts[2::2])}


def run_queries() -> dict[str, pd.DataFrame]:
    results = {}
    with sqlite3.connect(DB_PATH) as conn:
        for name, sql in load_queries(KPI_SQL).items():
            results[name] = pd.read_sql_query(sql, conn)
    return results


def budget_reallocation(campaigns: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    """Move the budget of 'Cut' campaigns to 'Scale' campaigns (in proportion to
    their current spend) and estimate the extra revenue.

    Moved money is assumed to earn only REALLOCATION_EFFICIENCY x the receiving
    campaign's current ROAS, because returns usually fall as budget grows.

    Brand-keyword campaigns are left out as receivers: they only show ads to
    people already searching the brand name, so extra budget can't buy much
    more traffic there.
    """
    cut = campaigns[campaigns["action"] == "Cut"]
    scale = campaigns[(campaigns["action"] == "Scale")
                      & ~campaigns["campaign_name"].str.contains("Brand")].copy()

    budget_freed = cut["spend"].sum()
    revenue_lost = cut["revenue"].sum()
    scale["extra_budget"] = budget_freed * scale["spend"] / scale["spend"].sum()
    scale["extra_revenue"] = scale["extra_budget"] * scale["roas"] * REALLOCATION_EFFICIENCY

    plan = pd.concat([
        cut.assign(budget_change=-cut["spend"], revenue_change=-cut["revenue"]),
        scale.assign(budget_change=scale["extra_budget"], revenue_change=scale["extra_revenue"]),
    ])[["platform", "campaign_name", "action", "roas", "budget_change", "revenue_change"]].round(2)

    summary = {
        "budget_moved": budget_freed,
        "revenue_lost": revenue_lost,
        "revenue_gained": scale["extra_revenue"].sum(),
        "net_revenue_change": scale["extra_revenue"].sum() - revenue_lost,
    }
    return plan, summary


def inr(x: float) -> str:
    """Format rupees in lakhs for readable insights (1 lakh = 1,00,000)."""
    return f"₹{x / 1e5:,.2f} lakh"


def write_insights(r: dict[str, pd.DataFrame], plan_summary: dict) -> list[str]:
    o = r["overall_kpis"].iloc[0]
    plat = r["platform_kpis"]
    camp = r["campaign_kpis"]
    best, worst = plat.iloc[0], plat.iloc[-1]
    wasters = r["budget_wasters"]
    scale = camp[camp["action"] == "Scale"]

    fat = r["creative_fatigue"]
    last = fat[fat["month"] == fat["month"].max()].sort_values("ctr_change_vs_jan_pct")
    tired = last.iloc[0]

    age = r["age_group_kpis"].sort_values("cpa")
    dev = r["device_kpis"].set_index("device")

    lines = [
        f"**Overall:** {inr(o.spend)} spent across 5 platforms returned {inr(o.revenue)} "
        f"in revenue - a blended ROAS of **{o.roas:.2f}x** from {int(o.conversions):,} conversions.",

        f"**Best platform:** {best.platform} returned **{best.roas:.2f}x** ROAS. "
        f"**Weakest:** {worst.platform} at **{worst.roas:.2f}x**, below the {TARGET_ROAS:.1f}x target.",

        f"**Money-losing campaigns:** {len(wasters)} campaign{'s' if len(wasters) != 1 else ''} "
        f"returned less than {'they' if len(wasters) != 1 else 'it'} cost "
        f"({', '.join(wasters['platform'] + ' - ' + wasters['campaign_name'])}), "
        f"losing {inr(-wasters['net_return'].sum())}.",

        f"**Under-funded winners:** {len(scale)} campaigns have ROAS of 4x or more but get only "
        f"{100 * scale['spend'].sum() / camp['spend'].sum():.0f}% of the budget "
        f"(best: {camp.iloc[0].platform} - {camp.iloc[0].campaign_name} at {camp.iloc[0].roas:.2f}x).",

        f"**Creative fatigue:** CTR on {tired.platform} - {tired.campaign_name} fell "
        f"**{abs(tired.ctr_change_vs_jan_pct):.0f}%** from January to June. The creative needs refreshing.",

        f"**Audience:** The {age.iloc[0].age_group} age group has the lowest cost per conversion "
        f"(₹{age.iloc[0].cpa:,.0f}); {age.iloc[-1].age_group} is the most expensive (₹{age.iloc[-1].cpa:,.0f}).",

        f"**Device:** Desktop converts at {dev.loc['Desktop', 'cvr_pct']:.2f}% vs "
        f"{dev.loc['Mobile', 'cvr_pct']:.2f}% on mobile, even though mobile gets "
        f"{100 * dev.loc['Mobile', 'spend'] / dev['spend'].sum():.0f}% of spend.",

        f"**Budget reallocation (estimate):** Moving {inr(plan_summary['budget_moved'])} from the 'Cut' "
        f"campaigns to the non-brand 'Scale' campaigns could add about "
        f"**{inr(plan_summary['net_revenue_change'])}** in revenue over 6 months at the same total spend.",
    ]

    md = ["# Key Insights", "", *[f"{i}. {line}" for i, line in enumerate(lines, 1)], "",
          "_Reallocation assumes moved budget earns "
          f"{int(REALLOCATION_EFFICIENCY * 100)}% of the receiving campaign's current ROAS. "
          "Brand-keyword search is excluded as a receiver because its traffic is limited "
          "by how many people search the brand name._", ""]
    (OUTPUT_DIR / "insights.md").write_text("\n".join(md), encoding="utf-8")
    return lines


def write_excel(sheets: dict[str, pd.DataFrame]):
    with pd.ExcelWriter(EXCEL_REPORT, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2A78D6")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for i, col in enumerate(df.columns, 1):
                letter = get_column_letter(i)
                width = max(len(str(col)), *(len(str(v)) for v in df[col].head(200))) + 2
                ws.column_dimensions[letter].width = min(width, 60)
                if col in ("spend", "revenue", "cpa", "cpc", "net_return", "budget_change", "revenue_change"):
                    for cell in ws[letter][1:]:
                        cell.number_format = "₹#,##0.00"


def main():
    OUTPUT_DIR.mkdir(exist_ok=True)
    results = run_queries()

    camp = results["campaign_kpis"]
    camp["action"] = camp["roas"].apply(roas_action)

    plan, plan_summary = budget_reallocation(camp)
    results["budget_reallocation"] = plan

    sql_dir = OUTPUT_DIR / "sql_results"
    sql_dir.mkdir(exist_ok=True)
    for name, df in results.items():
        df.to_csv(sql_dir / f"{name}.csv", index=False)

    insights = write_insights(results, plan_summary)

    write_excel({
        "Summary": results["overall_kpis"],
        "Platforms": results["platform_kpis"],
        "Campaigns": camp,
        "Budget Wasters": results["budget_wasters"],
        "Reallocation Plan": plan,
        "Monthly Trend": results["monthly_platform_trend"],
        "Creative Fatigue": results["creative_fatigue"],
        "Age Groups": results["age_group_kpis"],
        "Devices": results["device_kpis"],
        "Insights": pd.DataFrame({"insight": [re.sub(r"\*\*", "", s) for s in insights]}),
    })

    print(f"[4/5] Ran {len(results) - 1} SQL queries -> outputs/sql_results/, Excel report, insights")
    print()
    for line in insights:
        print("      - " + re.sub(r"\*\*", "", line))
    print()


if __name__ == "__main__":
    main()
